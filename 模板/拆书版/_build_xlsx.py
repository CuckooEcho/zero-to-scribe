#!/usr/bin/env python3
"""生成三张长篇小说管理表的 xlsx：章表 / 伏笔表 / 人物状态表（含凡人修仙传七玄门卷示例数据）"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT_DIR = "/Users/panxuan/Projects/feishu/03-hobby/zero-to-scribe/模板"

HEADER_FILL = PatternFill("solid", fgColor="1C1917")
HEADER_FONT = Font(name="PingFang SC", bold=True, color="FFFFFF", size=11)
BODY_FONT = Font(name="PingFang SC", size=11)
NOTE_FONT = Font(name="PingFang SC", size=10, color="78716C")
THIN = Side(style="thin", color="D6D3D1")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(vertical="top", wrap_text=True)
CENTER = Alignment(vertical="top", horizontal="center", wrap_text=True)


def build(path, title, headers, widths, rows, notes, center_cols=()):
    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws.sheet_view.showGridLines = False

    ws.append(headers)
    for c in ws[1]:
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = CENTER
        c.border = BORDER
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    for row in rows:
        ws.append(row)
    for r in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(headers)):
        for c in r:
            c.font = BODY_FONT
            c.border = BORDER
            c.alignment = CENTER if c.column in center_cols else WRAP

    ws.freeze_panes = "A2"

    # 填写规范 sheet
    ws2 = wb.create_sheet("填写规范")
    ws2.sheet_view.showGridLines = False
    ws2.column_dimensions["A"].width = 110
    ws2.append([f"{title} · 填写规范"])
    ws2["A1"].font = Font(name="PingFang SC", bold=True, size=13)
    for line in notes:
        ws2.append([line])
        ws2.cell(row=ws2.max_row, column=1).font = NOTE_FONT
        ws2.cell(row=ws2.max_row, column=1).alignment = WRAP

    wb.save(path)
    print("saved:", path)


# ---------- 章表 ----------
build(
    f"{OUT_DIR}/章表.xlsx",
    "章表",
    ["卷/章", "故事时间", "场景/出场", "本章目标", "阻碍", "结果（价值+/-）", "钩子", "状态",
     "情绪压强", "伏笔·埋", "伏笔·收", "主线进度", "字数"],
    [8, 10, 18, 22, 20, 26, 22, 8, 9, 9, 12, 12, 8],
    [
        ["1-1", "初春", "山边小村/韩立、父母", "家里太穷，想进七玄门挣钱", "选拔竞争", "入选记名弟子（+）", "七玄门里什么样？", "改定", 2, "", "", "凡人起步", ""],
        ["1-2", "数日后", "七玄门/韩立、同批弟子", "通过入门考核留下", "体能不过关", "被淘汰边缘被墨大夫挑走（±）", "墨大夫为什么专挑差的？", "改定", 3, "F02", "", "入七玄门", ""],
        ["1-3", "同日", "神手谷/墨大夫、张铁", "拜师学医", "墨大夫规矩古怪", "拜入墨大夫门下（+）", "谷里没有其他弟子", "改定", 3, "F02", "", "得庇护", ""],
        ["1-4", "数月", "神手谷/韩立", "练成长春功第一层", "功法进境极慢", "苦修入门，感到墨大夫态度渐冷（-）", "长春功到底有什么用？", "改定", 4, "F03", "", "接触修行", ""],
        ["1-5", "秋", "后山/韩立", "采药", "山崖危险", "捡到神秘小绿瓶（+）", "瓶子夜里吸月光", "改定", 5, "F01", "", "金手指入手", ""],
        ["1-6", "数周后", "神手谷/韩立", "弄清瓶子的用途", "不敢暴露", "发现绿液可催熟草药（+）", "这事绝不能让师父知道", "改定", 6, "F01", "", "金手指验证", ""],
        ["1-7", "年末", "七玄门/韩立、厉飞雨", "交朋友、站稳脚", "厉飞雨靠吃药硬撑武功", "以药相助结交厉飞雨（+）", "厉飞雨的瘾是隐患", "改定", 4, "F04", "", "第一个盟友", ""],
        ["1-8", "次年", "神手谷/墨大夫、韩立", "继续修炼兼学医", "墨大夫身体明显衰败、性情渐异", "韩立暗中警惕，开始藏拙（-）", "张铁不见了", "改定", 7, "F05", "", "危机逼近", ""],
        ["1-9", "数日后", "神手谷/韩立", "查清张铁下落", "墨大夫监视", "发现张铁被制成炼尸（-）", "下一个会不会是我？", "改定", 8, "", "F05", "敌人现形", ""],
        ["1-10", "当夜", "神手谷/韩立、墨大夫", "保住性命与神智", "墨大夫摊牌欲夺舍", "借长春功和暗算反杀，墨大夫死（±）", "墨大夫背后还有修仙者", "改定", 10, "F06", "F02、F03", "卷级高潮", ""],
    ],
    [
        "示例数据：《凡人修仙传》七玄门卷（按公开梗概整理的示意切分，正式拆书以原文为准）。",
        "",
        "本章目标：视角人物在这一章想要什么。填不出 → 这章可能该删。",
        "阻碍：谁/什么拦着。没有阻碍的一章是流水账。",
        "结果：目标达成 / 失败 / 反转，并标价值方向（+ 变好 / - 变坏 / ± 反转）。",
        "钩子：章末未闭合的悬念或危机。没有钩子读者就睡了。",
        "情绪压强：1–10，压抑/蓄力为低，爆发为高；拉成曲线看，连续十章低压会掉读者。",
        "伏笔列只写编号（如 F03），内容在《伏笔表.xlsx》维护，别在这里展开。",
        "状态：构思 → 初稿 → 改定。",
        "核心 8 列（至「状态」）写作时必填；审计列（情绪压强起）修订或卡文时再用。",
    ],
    center_cols=(1, 2, 8, 9, 10, 11, 12, 13),
)

# ---------- 伏笔表 ----------
build(
    f"{OUT_DIR}/伏笔表.xlsx",
    "伏笔表",
    ["编号", "埋设章", "内容（一句话）", "埋设方式", "计划回收章", "实际回收章", "状态"],
    [8, 10, 40, 16, 12, 12, 22],
    [
        ["F01", "1-5", "小绿瓶催熟灵药，来历不明", "道具", "全书终章级", "", "未回收（终极悬念）"],
        ["F02", "1-2", "墨大夫专挑被淘汰的弟子", "细节（反常）", "1-10", "1-10", "已回收"],
        ["F03", "1-4", "长春功进境慢却要求必须练", "细节（反常）", "1-10", "1-10", "已回收"],
        ["F04", "1-7", "厉飞雨靠药物硬撑，身体有隐患", "人物", "本卷末", "", "未回收（卷级钩子）"],
        ["F05", "1-8", "张铁失踪", "细节", "1-9", "1-9", "已回收"],
        ["F06", "1-10", "墨大夫背后有真正的修仙界", "对话", "下一卷", "", "未回收（升维钩子）"],
    ],
    [
        "一条伏笔一行：生命周期横跨几十章，独立成表，不塞进章表。",
        "",
        "埋设方式：道具（给了个东西）/ 对话（提到个名字）/ 细节（反常描写）/ 人物（登场未展开）。",
        "状态：未回收 → 已回收 → 废弃（废弃也要记，避免读者记得你忘了）。",
        "修订用法：筛「未回收」且埋设章距今超过两卷的——要么赶紧收，要么标记废弃。",
        "契诃夫原则：挂墙上的枪必须响；不响的枪别挂。",
        "健康节奏参考：卷内埋卷内收约占一半（闭合感），留一半到卷外（续命感）。",
    ],
    center_cols=(1, 2, 5, 6),
)

# ---------- 人物状态表 ----------
wb = Workbook()
ws = wb.active
ws.title = "主角快照"
ws.sheet_view.showGridLines = False
snap_headers = ["卷", "境界/等级", "核心资源", "持有物", "伤势/隐患", "主要仇家", "关键关系变化", "当前目标"]
snap_widths = [12, 18, 22, 26, 24, 26, 18, 20]
ws.append(snap_headers)
for c in ws[1]:
    c.fill, c.font, c.alignment, c.border = HEADER_FILL, HEADER_FONT, CENTER, BORDER
for i, w in enumerate(snap_widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.append(["七玄门卷末", "长春功数层（刚入门）", "催熟草药的隐秘渠道", "小绿瓶、墨大夫遗留的修仙者物品",
           "无（但身怀秘密=最大隐患）", "暂无（墨大夫已死，其背后势力未知）", "厉飞雨成生死之交", "弄清修仙界，活下去"])
for r in ws.iter_rows(min_row=2, max_col=len(snap_headers)):
    for c in r:
        c.font, c.border, c.alignment = BODY_FONT, BORDER, WRAP
ws.freeze_panes = "A2"

ws2 = wb.create_sheet("配角档案")
ws2.sheet_view.showGridLines = False
npc_headers = ["人物", "首登场章", "身份/阵营", "与主角关系", "欲望（ta 想要什么）", "秘密（ta 瞒着什么）", "状态"]
npc_widths = [10, 10, 22, 14, 26, 28, 16]
ws2.append(npc_headers)
for c in ws2[1]:
    c.fill, c.font, c.alignment, c.border = HEADER_FILL, HEADER_FONT, CENTER, BORDER
for i, w in enumerate(npc_widths, 1):
    ws2.column_dimensions[get_column_letter(i)].width = w
for row in [
    ["墨大夫", "1-2", "七玄门供奉/落魄修仙者", "师父→敌人", "夺舍重生，重返修仙界", "收徒是为挑选夺舍躯壳", "死亡"],
    ["张铁", "1-3", "同门师兄", "同门", "习武出头", "（不知情）", "死亡（被炼成尸）"],
    ["厉飞雨", "1-7", "七玄门弟子", "生死之交", "成为高手", "靠药物透支身体", "存活（隐患未解）"],
]:
    ws2.append(row)
for r in ws2.iter_rows(min_row=2, max_col=len(npc_headers)):
    for c in r:
        c.font, c.border, c.alignment = BODY_FONT, BORDER, (CENTER if c.column in (1, 2, 7) else WRAP)
ws2.freeze_panes = "A2"

ws3 = wb.create_sheet("填写规范")
ws3.sheet_view.showGridLines = False
ws3.column_dimensions["A"].width = 110
ws3.append(["人物状态表 · 填写规范"])
ws3["A1"].font = Font(name="PingFang SC", bold=True, size=13)
for line in [
    "示例数据：《凡人修仙传》七玄门卷（按公开梗概整理，非逐章原文校对）。",
    "",
    "跨章的持久状态（roguelike 里的角色 build），每卷末更新一次快照，不要逐章重复填。",
    "章表记「发生了什么」，这张表记「此刻是什么样」。",
    "持有物：只记有剧情功能的（金手指、关键道具），行囊流水账不记。",
    "欲望和秘密是配角的两根发条：没有这两项的配角是工具人。",
    "状态：存活 / 死亡 / 失联 / 叛变——失联的人也是伏笔。",
    "卷末「隐患」栏直接生成下一卷的开局矛盾（分形三幕的接龙机制）。",
]:
    ws3.append([line])
    ws3.cell(row=ws3.max_row, column=1).font = NOTE_FONT
    ws3.cell(row=ws3.max_row, column=1).alignment = WRAP

path = f"{OUT_DIR}/人物状态表.xlsx"
wb.save(path)
print("saved:", path)
