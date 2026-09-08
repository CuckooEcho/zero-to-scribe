#!/usr/bin/env python3
"""生成写书版两张表：事件表 / 场景表（含七玄门卷反向演示数据）"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT_DIR = "/Users/panxuan/Projects/feishu/03-hobby/zero-to-scribe/模板/写书版"

HEADER_FILL = PatternFill("solid", fgColor="1C1917")
HEADER_FONT = Font(name="PingFang SC", bold=True, color="FFFFFF", size=11)
BODY_FONT = Font(name="PingFang SC", size=11)
NOTE_FONT = Font(name="PingFang SC", size=10, color="78716C")
THIN = Side(style="thin", color="D6D3D1")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(vertical="top", wrap_text=True)
CENTER = Alignment(vertical="top", horizontal="center", wrap_text=True)


def build(path, title, headers, widths, rows, notes, center_cols=()):
    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws.sheet_view.showGridLines = False
    ws.append(headers)
    for c in ws[1]:
        c.fill, c.font, c.alignment, c.border = HEADER_FILL, HEADER_FONT, CENTER, BORDER
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in rows:
        ws.append(row)
    for r in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(headers)):
        for c in r:
            c.font = BODY_FONT
            c.border = BORDER
            c.alignment = CENTER if c.column in center_cols else WRAP
    ws.freeze_panes = "A2"

    ws2 = wb.create_sheet("填写规范")
    ws2.sheet_view.showGridLines = False
    ws2.column_dimensions["A"].width = 110
    ws2.append([f"{title} · 填写规范"])
    ws2["A1"].font = Font(name="PingFang SC", bold=True, size=13)
    for line in notes:
        ws2.append([line])
        ws2.cell(row=ws2.max_row, column=1).font = NOTE_FONT
        ws2.cell(row=ws2.max_row, column=1).alignment = WRAP

    wb.save(path)
    print("saved:", path)


# ---------- 事件表 ----------
build(
    f"{OUT_DIR}/事件表.xlsx",
    "事件表",
    ["事件编号", "事件名", "三幕位置", "目标（谁想要什么）", "冲突（对手/阻碍）", "结果", "价值转变", "场景数", "卷尾钩子贡献"],
    [9, 12, 14, 26, 22, 30, 20, 8, 28],
    [
        ["E01", "入门", "开场+激励事件", "韩立想进七玄门脱贫", "考核不合格", "被墨大夫挑中收徒", "±（脱贫了，但入了个局）", 3, "墨大夫的挑选标准可疑"],
        ["E02", "得宝", "中点前", "想在门内站稳脚跟", "功法进境慢、被冷落", "后山捡到小绿瓶，验证可催熟灵药", "+（拿到金手指）", 3, "瓶子来历（全书级悬念 F01）"],
        ["E03", "结盟", "中点", "需要有盟友互相照应", "厉飞雨药物透支", "以药相交，结下生死之交", "+", 1, "厉飞雨的隐患（F04）"],
        ["E04", "夺舍危机", "低谷+高潮", "保住性命与神智", "墨大夫摊牌夺舍", "反杀成功", "±（活下来，但师父死了、秘密上身）", 3, "墨大夫背后的修仙界（F06）"],
    ],
    [
        "示例数据：用《凡人修仙传》七玄门卷反向演示写作流程（示意切分，非原文校对）。",
        "",
        "每行一个事件：一次完整的「目标→对抗→结果」，价值发生一次大转变。一卷 5–8 个事件为宜。",
        "三幕位置：本卷内部的定位——开场 / 激励事件 / 中点转折 / 低谷 / 高潮 / 收束。",
        "价值转变：+ 向好 / - 向坏 / ± 反转。没有价值转变的不是事件，是过渡段。",
        "自查：相邻事件的价值转变应当正负交替——连续三个 + 太顺，连续三个 - 太丧。",
        "卷尾钩子贡献：本事件为下一卷留下了什么（新敌人 / 新悬念 / 新目标）。",
        "配套：场景表.xlsx（展开期）；伏笔表与人物状态表见 ../拆书版/。",
    ],
    center_cols=(1, 7, 8),
)

# ---------- 场景表 ----------
build(
    f"{OUT_DIR}/场景表.xlsx",
    "场景表",
    ["事件/场景", "故事时间", "场景/出场", "本场景目标", "阻碍", "结果（+/-）", "情绪压强", "伏笔·埋", "伏笔·收", "钩子", "建议切章点", "状态", "字数"],
    [10, 10, 18, 20, 22, 24, 9, 9, 12, 22, 11, 8, 8],
    [
        ["E04-S1", "次年", "神手谷/墨大夫、韩立", "继续修炼同时藏拙", "墨大夫身体衰败、性情渐异", "韩立起戒心（-）", 7, "F05", "", "张铁不见了", "切", "构思", ""],
        ["E04-S2", "数日后", "神手谷/韩立", "查清张铁下落", "墨大夫监视", "发现张铁被制成炼尸（-）", 8, "", "F05", "下一个会不会是我？", "切", "构思", ""],
        ["E04-S3", "当夜", "神手谷/韩立、墨大夫", "反杀", "实力悬殊", "墨大夫死（±）", 10, "F06", "F02、F03", "背后还有修仙者", "", "构思", ""],
    ],
    [
        "示例数据：仅展开 E04「夺舍危机」演示跨章事件的表示法（共享事件编号，筛选 E04 即见完整弧线）。",
        "",
        "每行一个场景：一段连续时空里的戏。编号 事件-场景（如 E04-S2）。",
        "此时不写章号——章是包装期看着「建议切章点」列才切出来的。",
        "本场景目标：事件目标的子任务，不是复述。",
        "结果：场景的微型价值转变，结束时处境必须和开始时不同。",
        "情绪压强：1–10，卷级高潮场景应到 9–10。",
        "建议切章点：标「切」= 钩子最痒、可在此断章；高潮场景后面要跟收束段，不宜立刻断。",
        "写正文时同一事件的场景连续写，保持情绪连贯；切章最后再做。",
        "伏笔列只写编号，内容在 ../拆书版/伏笔表.xlsx 维护。",
    ],
    center_cols=(1, 2, 7, 8, 9, 11, 12, 13),
)
